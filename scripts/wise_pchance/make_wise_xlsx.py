#!/usr/bin/env python
"""outputs/wise_pchance.csv -> xlsx, in the same styling as
outputs/gaia_pchance_galactic*.xlsx (Arial 10, white-on-#3B6FB6 header,
freeze B2, autofilter, zebra rows, per-column decimal formats, tiny p in
scientific notation).

Sheet 1 "wise_pchance"  the 73 sources, sorted by p (primary floor first)
Sheet 2 "wise_vs_gaia"  the same p beside the Galactic-field Gaia p

Run: python make_wise_xlsx.py
"""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CSV = os.path.join(REPO, "outputs", "wise_pchance.csv")
CMP = os.path.join(REPO, "outputs", "wise_vs_gaia_pchance.csv")
OUT = os.path.join(REPO, "outputs", "wise_pchance.xlsx")

P_FMT = "[<0.001]0.00E+00;0.0000"   # tiny p in scientific, rest 4 decimals
MAIN = ["spt_id", "p_wise_gal", "p_wise_winter", "association_TS_gal",
        "association_TS_winter", "event_ts", "mapping_ts", "n_null_ge_gal",
        "sep_arcsec", "wise_designation", "w1mpro", "w2mpro", "w1_rank",
        "n_cone", "n_wise_10arcsec", "cc_flags", "ext_flg", "ph_qual",
        "ra_deg", "dec_deg", "snr_max", "wise_ra", "wise_dec",
        "sigma_ra_gal", "sigma_dec_gal", "sigma_ra_winter", "sigma_dec_winter"]
COMPARE = ["spt_id", "p_wise_gal", "p_gaia_gal", "association_TS_gal",
           "association_TS_gaia", "sep_arcsec", "gaia_sep_arcsec", "w1mpro",
           "event_ts", "mapping_ts"]

FMT = {
    "p_wise_gal": P_FMT, "p_wise_winter": P_FMT, "p_gaia_gal": P_FMT,
    "association_TS_gal": "0.00", "association_TS_winter": "0.00",
    "association_TS_gaia": "0.00", "event_ts": "0.0", "mapping_ts": "0",
    "n_null_ge_gal": "0", "sep_arcsec": "0.00", "gaia_sep_arcsec": "0.00",
    "wise_designation": "@", "w1mpro": "0.000", "w2mpro": "0.000",
    "w1_rank": "0", "n_cone": "0", "n_wise_10arcsec": "0", "cc_flags": "@",
    "ext_flg": "0", "ph_qual": "@", "ra_deg": "0.00000", "dec_deg": "0.00000",
    "snr_max": "0.0", "wise_ra": "0.00000", "wise_dec": "0.00000",
    "sigma_ra_gal": "0.00", "sigma_dec_gal": "0.00",
    "sigma_ra_winter": "0.00", "sigma_dec_winter": "0.00",
}
WIDTHS = {"spt_id": 26, "wise_designation": 21, "association_TS_winter": 15,
          "association_TS_gal": 15, "association_TS_gaia": 16,
          "n_wise_10arcsec": 15, "n_null_ge_gal": 13, "gaia_sep_arcsec": 14,
          "sigma_dec_winter": 15, "sigma_ra_winter": 14}
TEXT_COLS = {"wise_designation", "cc_flags", "ph_qual", "spt_id"}

HDR_FONT = Font(name="Arial", size=10, bold=True, color="00FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="003B6FB6")
BODY_FONT = Font(name="Arial", size=10)
ZEBRA = [PatternFill("solid", fgColor="00DCE9F7"),
         PatternFill("solid", fgColor="00FCE8D4")]


def write_sheet(ws, df, order):
    for j, col in enumerate(order, 1):
        c = ws.cell(row=1, column=j, value=col)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(col, 13)
    for i, row in enumerate(df[order].itertuples(index=False), 2):
        fill = ZEBRA[i % 2]
        for j, (col, v) in enumerate(zip(order, row), 1):
            if col in TEXT_COLS:
                v = str(v)
            elif pd.isna(v):
                v = None
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.fill = fill
            c.number_format = FMT.get(col, "General")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(order))}{len(df) + 1}"


def main():
    df = pd.read_csv(CSV).sort_values("p_wise_gal").reset_index(drop=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "wise_pchance"
    write_sheet(ws, df, MAIN)

    if os.path.isfile(CMP):
        cmp = pd.read_csv(CMP).sort_values("p_wise_gal").reset_index(drop=True)
        write_sheet(wb.create_sheet("wise_vs_gaia"), cmp, COMPARE)

    wb.save(OUT)
    print(f"wrote {OUT}: {len(df)} rows x {len(MAIN)} cols"
          + (f" + wise_vs_gaia {len(cmp)} rows" if os.path.isfile(CMP) else ""))


if __name__ == "__main__":
    main()
