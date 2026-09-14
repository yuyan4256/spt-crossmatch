#!/usr/bin/env python
"""outputs/gaia_pchance_galactic.csv -> xlsx, mirroring gaia_10arcsec(1).xlsx
styling (Arial 10, white-on-#3B6FB6 header, freeze B2, autofilter, source_id
as text, per-column decimal formats). p_galactic sits beside p_chance."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CSV = "outputs/gaia_pchance_galactic.csv"
OUT = "outputs/gaia_pchance_galactic.xlsx"

ORDER = ["spt_id", "p_chance", "p_galactic", "event_ts", "mapping_ts",
         "spt_ra_deg", "spt_dec_deg", "spt_snr_max", "source_id", "ra", "dec",
         "gaia_sep_arcsec", "phot_g_mean_mag", "association_TS", "parallax",
         "parallax_over_error", "pmra", "pmra_error", "pmdec", "pmdec_error"]
P_FMT = "[<0.001]0.00E+00;0.0000"  # tiny p in scientific, rest 4 decimals
FMT = {"p_chance": P_FMT, "p_galactic": P_FMT, "event_ts": "0.0",
       "mapping_ts": "0", "spt_ra_deg": "0.00000", "spt_dec_deg": "0.00000",
       "spt_snr_max": "0.0", "source_id": "@", "ra": "0.00000",
       "dec": "0.00000", "gaia_sep_arcsec": "0.00", "phot_g_mean_mag": "0.00",
       "association_TS": "0.00", "parallax": "0.000",
       "parallax_over_error": "0.00", "pmra": "0.000", "pmra_error": "0.000",
       "pmdec": "0.000", "pmdec_error": "0.000"}
WIDTHS = {"spt_id": 26, "source_id": 21}

df = pd.read_csv(CSV)[ORDER]

wb = Workbook()
ws = wb.active
ws.title = "gaia_pchance_galactic"

hdr_font = Font(name="Arial", size=10, bold=True, color="00FFFFFF")
hdr_fill = PatternFill("solid", fgColor="003B6FB6")
body_font = Font(name="Arial", size=10)
# zebra striping per spt_id block, as in gaia_10arcsec(1).xlsx
zebra = [PatternFill("solid", fgColor="00DCE9F7"),
         PatternFill("solid", fgColor="00FCE8D4")]

for j, col in enumerate(ORDER, 1):
    c = ws.cell(row=1, column=j, value=col)
    c.font = hdr_font
    c.fill = hdr_fill
    ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(col, 13)

block, prev_id = -1, None
for i, row in enumerate(df.itertuples(index=False), 2):
    if row.spt_id != prev_id:
        block, prev_id = block + 1, row.spt_id
    fill = zebra[block % 2]
    for j, (col, v) in enumerate(zip(ORDER, row), 1):
        if col == "source_id":
            v = str(v)
        elif pd.isna(v):
            v = None
        c = ws.cell(row=i, column=j, value=v)
        c.font = body_font
        c.fill = fill
        c.number_format = FMT.get(col, "General")

ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(ORDER))}{len(df) + 1}"
wb.save(OUT)
print(f"wrote {OUT}: {len(df)} rows x {len(ORDER)} cols")
