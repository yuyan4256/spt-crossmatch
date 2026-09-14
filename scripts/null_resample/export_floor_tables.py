#!/usr/bin/env python
"""Split outputs/gaia_pchance_floors.csv into one CSV + one XLSX per pointing
floor, named after that floor, in the same layout/styling as
"outputs/gaia_pchance_galactic 3.05'' 4.42''.xlsx".

  gaia_pchance_galactic 3.62'' 4.85''.csv / .xlsx     abs (+) mean sigma_sys,rel
  gaia_pchance_galactic 3.78'' 4.87''.csv / .xlsx     abs (+) max  sigma_sys,rel

Column meaning in the per-floor files (identical layout to the 3.05/4.42 one):
  association_TS  this floor's TS, recomputed on the local Gaia catalog
  p_galactic      this floor's p_chance, from the matching null curve
  p_chance        untouched: the shipped winter-table p, kept for reference
"""
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
from run_floors_rerun import FLOORS

SRC = os.path.join(REPO, "outputs", "gaia_pchance_floors.csv")
OUTDIR = os.path.join(REPO, "outputs")

ORDER = ["spt_id", "p_chance", "p_galactic", "event_ts", "mapping_ts",
         "spt_ra_deg", "spt_dec_deg", "spt_snr_max", "source_id", "ra", "dec",
         "gaia_sep_arcsec", "phot_g_mean_mag", "association_TS", "parallax",
         "parallax_over_error", "pmra", "pmra_error", "pmdec", "pmdec_error"]
P_FMT = "[<0.001]0.00E+00;0.0000"
FMT = {"p_chance": P_FMT, "p_galactic": P_FMT, "event_ts": "0.0",
       "mapping_ts": "0", "spt_ra_deg": "0.00000", "spt_dec_deg": "0.00000",
       "spt_snr_max": "0.0", "source_id": "@", "ra": "0.00000",
       "dec": "0.00000", "gaia_sep_arcsec": "0.00", "phot_g_mean_mag": "0.00",
       "association_TS": "0.00", "parallax": "0.000",
       "parallax_over_error": "0.00", "pmra": "0.000", "pmra_error": "0.000",
       "pmdec": "0.000", "pmdec_error": "0.000"}
WIDTHS = {"spt_id": 26, "source_id": 21}


def write_xlsx(df, path, sheet):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    hdr_font = Font(name="Arial", size=10, bold=True, color="00FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="003B6FB6")
    body_font = Font(name="Arial", size=10)
    zebra = [PatternFill("solid", fgColor="00DCE9F7"),
             PatternFill("solid", fgColor="00FCE8D4")]

    for j, col in enumerate(ORDER, 1):
        c = ws.cell(row=1, column=j, value=col)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(col, 13)

    block, prev = -1, None
    for i, row in enumerate(df[ORDER].itertuples(index=False), 2):
        if row.spt_id != prev:
            block, prev = block + 1, row.spt_id
        fill = zebra[block % 2]
        for j, (col, v) in enumerate(zip(ORDER, row), 1):
            if col == "source_id":
                v = str(v)
            elif pd.isna(v):
                v = None
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.fill = body_font, fill
            c.number_format = FMT.get(col, "General")

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ORDER))}{len(df) + 1}"
    wb.save(path)


def main():
    src = pd.read_csv(SRC)
    for name, (fra, fdec) in FLOORS.items():
        if name.startswith("F0"):
            continue          # 3.05/4.42 already shipped as the production file
        df = src.copy()
        df["association_TS"] = df[f"association_TS_{name}"]
        df["p_galactic"] = df[f"p_{name}"]
        stem = f"gaia_pchance_galactic {fra:.2f}'' {fdec:.2f}''"
        csv = os.path.join(OUTDIR, stem + ".csv")
        xlsx = os.path.join(OUTDIR, stem + ".xlsx")
        df[ORDER].to_csv(csv, index=False)
        write_xlsx(df, xlsx, stem)
        n = df.loc[df.groupby("spt_id")["association_TS"].idxmax()]
        print(f"{name:18s} floor {fra:.3f}\"/{fdec:.3f}\"  "
              f"p<0.02: {int((n.p_galactic < 0.02).sum())}  -> {stem}.csv/.xlsx")


if __name__ == "__main__":
    main()
